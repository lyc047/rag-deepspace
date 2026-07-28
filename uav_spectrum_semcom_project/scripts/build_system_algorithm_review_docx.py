#!/usr/bin/env python
"""Build the audit-ready Word handoff from the Markdown source."""

from __future__ import annotations

import re
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT = Path(__file__).resolve().parents[1]
SOURCE = PROJECT / "docs" / "SYSTEM_ALGORITHM_REVIEW_HANDOFF_20260727.md"
OUTPUT = PROJECT / "docs" / "SYSTEM_ALGORITHM_REVIEW_HANDOFF_20260727.docx"
ASSETS = PROJECT / "docs" / "review_handoff_assets"
ARCHITECTURE = ASSETS / "current_architecture_flow.png"
EVOLUTION = ASSETS / "algorithm_evolution.png"

BLUE = "1F4E79"
MID_BLUE = "5B9BD5"
LIGHT_BLUE = "DCE6F1"
PALE_BLUE = "EDF3F8"
DARK = "243746"
GREY = "667788"
LIGHT_GREY = "F2F4F6"
WHITE = "FFFFFF"
RED = "A61B1B"
GREEN = "276749"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def create_numbering_instance(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    num_id = max(num_ids, default=0) + 1

    # Reuse Word's built-in "List Number" abstract definition.  Creating custom
    # abstractNum records is needlessly fragile across Word/LibreOffice versions;
    # a fresh num instance with startOverride is the schema-native way to restart.
    abstract_id = None
    for abstract in numbering.findall(qn("w:abstractNum")):
        for level in abstract.findall(qn("w:lvl")):
            paragraph_style = level.find(qn("w:pStyle"))
            if (
                paragraph_style is not None
                and paragraph_style.get(qn("w:val")) == "ListNumber"
            ):
                abstract_id = int(abstract.get(qn("w:abstractNumId")))
                break
        if abstract_id is not None:
            break
    if abstract_id is None:
        raise RuntimeError("Word built-in List Number definition was not found")

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return num_id


def add_numbered_item(doc: Document, text: str, display_number: int) -> None:
    # Use an explicit visible number for deterministic restart and cross-renderer
    # compatibility. Other unordered lists in the report retain native Word list
    # semantics; these ordered sequences are deliberately presentation-stable.
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.2)
    prefix = p.add_run(f"{display_number}. ")
    prefix.bold = True
    add_inline(p, text)


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_east_asia_font(run, font_name: str) -> None:
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")


def set_style_font(style, east_asia: str, latin: str = "Calibri") -> None:
    style.font.name = latin
    r_pr = style._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.32)

    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.widow_control = True

    for name, size, color, before, after in (
        ("Title", 25, BLUE, 0, 8),
        ("Subtitle", 13, GREY, 0, 8),
        ("Heading 1", 16, BLUE, 12, 7),
        ("Heading 2", 13, BLUE, 10, 5),
        ("Heading 3", 11.5, DARK, 8, 4),
    ):
        style = styles[name]
        set_style_font(style, "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        if name.startswith("Heading") or name == "Title":
            style.font.bold = True

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        set_style_font(style, "Microsoft YaHei")
        style.font.size = Pt(10.25)
        style.paragraph_format.left_indent = Inches(0.28)
        style.paragraph_format.first_line_indent = Inches(-0.18)
        style.paragraph_format.space_after = Pt(3)

    quote = styles["Quote"]
    set_style_font(quote, "Microsoft YaHei")
    quote.font.size = Pt(10.25)
    quote.font.color.rgb = RGBColor.from_string(DARK)
    quote.paragraph_format.left_indent = Inches(0.22)
    quote.paragraph_format.right_indent = Inches(0.10)
    quote.paragraph_format.space_before = Pt(5)
    quote.paragraph_format.space_after = Pt(7)

    if "Formula Readable" not in styles:
        formula = styles.add_style("Formula Readable", WD_STYLE_TYPE.PARAGRAPH)
        formula.base_style = styles["Normal"]
        set_style_font(formula, "Microsoft YaHei", "Cambria")
        formula.font.size = Pt(10.5)
        formula.font.color.rgb = RGBColor.from_string(BLUE)
        formula.paragraph_format.left_indent = Inches(0.35)
        formula.paragraph_format.right_indent = Inches(0.18)
        formula.paragraph_format.space_before = Pt(4)
        formula.paragraph_format.space_after = Pt(6)
        formula.paragraph_format.keep_together = True

    if "Small Note" not in styles:
        note = styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
        note.base_style = styles["Normal"]
        set_style_font(note, "Microsoft YaHei")
        note.font.size = Pt(8.5)
        note.font.color.rgb = RGBColor.from_string(GREY)
        note.paragraph_format.space_after = Pt(3)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("频谱资源语义算法 · 第三方审查稿")
    set_east_asia_font(run, "Microsoft YaHei")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GREY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("冻结状态：S6.6d开发验收完成；阶段6 Final访问0次  ·  ")
    set_east_asia_font(run, "Microsoft YaHei")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GREY)
    add_page_number(footer)


def font_prop(size: int, weight: str = "normal") -> FontProperties:
    candidates = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return FontProperties(fname=str(candidate), size=size, weight=weight)
    return FontProperties(size=size, weight=weight)


def draw_evolution() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(12, 3.0), dpi=180)
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 3)
    ax.axis("off")
    boxes = [
        (0.25, "阶段4\n任务充分块语义\n确定“传什么”", MID_BLUE),
        (3.25, "阶段5\n事件触发与多查询\n确定“何时传”", "4F81BD"),
        (6.25, "阶段6\n任务等价码本\n压缩“精确动作”", BLUE),
        (9.25, "可靠闭环\n信念·心跳·预测保护\n保证“安全执行”", "244062"),
    ]
    for x, text, color in boxes:
        patch = FancyBboxPatch(
            (x, 0.65),
            2.45,
            1.65,
            boxstyle="round,pad=0.08,rounding_size=0.12",
            facecolor=f"#{color}",
            edgecolor="white",
            linewidth=1.5,
        )
        ax.add_patch(patch)
        ax.text(
            x + 1.225,
            1.475,
            text,
            ha="center",
            va="center",
            color="white",
            fontproperties=font_prop(11, "bold"),
            linespacing=1.45,
        )
    for x1, x2 in ((2.70, 3.22), (5.70, 6.22), (8.70, 9.22)):
        ax.add_patch(
            FancyArrowPatch(
                (x1, 1.48),
                (x2, 1.48),
                arrowstyle="-|>",
                mutation_scale=18,
                color=f"#{GREY}",
                linewidth=1.6,
            )
        )
    fig.tight_layout(pad=0.25)
    fig.savefig(EVOLUTION, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def draw_architecture() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(12, 6.1), dpi=180)
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 7)
    ax.axis("off")
    nodes = [
        (0.3, 4.55, 2.0, 1.15, "真实功率扫频\nSigMF → N维功率", "DCE6F1", BLUE),
        (2.75, 4.55, 2.05, 1.15, "多查询任务状态\n代价·regret·年龄", "DCE6F1", BLUE),
        (5.25, 4.55, 2.1, 1.15, "硬触发与信念\n最坏状态判定", "FFF2CC", "806000"),
        (7.85, 4.55, 1.8, 1.15, "任务等价码本\n2-bit或escape", "E2F0D9", GREEN),
        (10.05, 4.55, 1.65, 1.15, "版本化更新\n身份校验", "E2F0D9", GREEN),
        (10.05, 2.35, 1.65, 1.15, "接收端执行\nfail-closed", "FCE4D6", RED),
        (7.65, 2.35, 2.0, 1.15, "ACK与上下文\n累计纪元过滤", "EDEDED", DARK),
        (5.1, 2.35, 2.05, 1.15, "固定20景心跳\n静默重置发现", "EDEDED", DARK),
        (2.6, 2.35, 2.05, 1.15, "一步风险排序\n只分配重复保护", "E4DFEC", "60497A"),
        (0.3, 2.35, 1.8, 1.15, "统一评价\nbit·clean·CVaR", "D9EAD3", GREEN),
    ]
    for x, y, w, h, text, fill, edge in nodes:
        patch = FancyBboxPatch(
            (x, y),
            w,
            h,
            boxstyle="round,pad=0.05,rounding_size=0.08",
            facecolor=f"#{fill}",
            edgecolor=f"#{edge}",
            linewidth=1.4,
        )
        ax.add_patch(patch)
        ax.text(
            x + w / 2,
            y + h / 2,
            text,
            ha="center",
            va="center",
            color=f"#{DARK}",
            fontproperties=font_prop(10, "bold"),
            linespacing=1.45,
        )
    arrows = [
        ((2.3, 5.13), (2.75, 5.13)),
        ((4.80, 5.13), (5.25, 5.13)),
        ((7.35, 5.13), (7.85, 5.13)),
        ((9.65, 5.13), (10.05, 5.13)),
        ((10.88, 4.55), (10.88, 3.50)),
        ((10.05, 2.93), (9.65, 2.93)),
        ((7.65, 2.93), (7.15, 2.93)),
        ((5.10, 2.93), (4.65, 2.93)),
        ((2.60, 2.93), (2.10, 2.93)),
        ((1.20, 3.50), (1.20, 4.55)),
        ((6.1, 3.50), (6.1, 4.55)),
        ((3.63, 3.50), (6.05, 4.50)),
    ]
    for start, end in arrows:
        ax.add_patch(
            FancyArrowPatch(
                start,
                end,
                arrowstyle="-|>",
                mutation_scale=14,
                color=f"#{GREY}",
                linewidth=1.3,
                connectionstyle="arc3,rad=0.0",
            )
        )
    ax.text(
        6.0,
        6.42,
        "当前冻结架构：预测不能选择未来频谱动作，也不能取消regret/年龄硬触发",
        ha="center",
        va="center",
        fontproperties=font_prop(12, "bold"),
        color=f"#{BLUE}",
    )
    ax.text(
        6.0,
        0.78,
        "数据流从左上到右下；反馈与上下文恢复形成闭环。任何码本身份异常均拒绝执行。",
        ha="center",
        va="center",
        fontproperties=font_prop(9),
        color=f"#{GREY}",
    )
    fig.tight_layout(pad=0.25)
    fig.savefig(ARCHITECTURE, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def add_inline(paragraph, text: str, *, default_bold=False, default_italic=False) -> None:
    token_re = re.compile(r"(\*\*.*?\*\*|`.*?`)")
    parts = token_re.split(text)
    for part in parts:
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        is_code = part.startswith("`") and part.endswith("`")
        content = part[2:-2] if is_bold else part[1:-1] if is_code else part
        run = paragraph.add_run(content)
        run.bold = default_bold or is_bold
        run.italic = default_italic
        set_east_asia_font(run, "Microsoft YaHei")
        if is_code:
            run.font.name = "Consolas"
            run.font.size = Pt(8.7)
            run.font.color.rgb = RGBColor.from_string(DARK)
            r_pr = run._element.get_or_add_rPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), LIGHT_GREY)
            r_pr.append(shd)


def add_cover(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    set_cell_width(table.cell(0, 0), 9300)
    cell = table.cell(0, 0)
    set_cell_shading(cell, BLUE)
    set_cell_margins(cell, top=260, start=280, bottom=260, end=280)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("RESEARCH HANDOFF  ·  TECHNICAL REVIEW")
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(WHITE)
    set_east_asia_font(run, "Microsoft YaHei")

    doc.add_paragraph()
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_inline(title, "基于离散任务导向语义通信的\n频谱资源选择系统", default_bold=True)
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_inline(subtitle, "完整算法、实现与证据审查稿")

    line = doc.add_table(rows=1, cols=1)
    line.autofit = False
    set_cell_width(line.cell(0, 0), 9300)
    set_cell_shading(line.cell(0, 0), MID_BLUE)
    set_cell_margins(line.cell(0, 0), top=35, start=0, bottom=35, end=0)

    doc.add_paragraph()
    status = doc.add_table(rows=4, cols=2)
    status.autofit = False
    labels = [
        ("文档版本", "2026-07-27"),
        ("冻结状态", "S6.6d开发验收完成；S6.7尚未开始"),
        ("外部Final", "阶段6访问次数为0，信号值未访问"),
        ("用途", "交由第三方审查创新性、可行性、证据强度与后续方向"),
    ]
    for i, (label, value) in enumerate(labels):
        set_cell_width(status.cell(i, 0), 1850)
        set_cell_width(status.cell(i, 1), 7450)
        set_cell_shading(status.cell(i, 0), LIGHT_BLUE)
        set_cell_shading(status.cell(i, 1), WHITE)
        for c in status.rows[i].cells:
            set_cell_margins(c, top=95, start=120, bottom=95, end=120)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p1 = status.cell(i, 0).paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        add_inline(p1, label, default_bold=True)
        p2 = status.cell(i, 1).paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        add_inline(p2, value)

    doc.add_paragraph()
    callout = doc.add_table(rows=1, cols=1)
    callout.autofit = False
    set_cell_width(callout.cell(0, 0), 9300)
    set_cell_shading(callout.cell(0, 0), PALE_BLUE)
    set_cell_margins(callout.cell(0, 0), top=150, start=180, bottom=150, end=180)
    p = callout.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline(
        p,
        "核心判定：当前候选已达到“可以注册外部Final协议”的开发成熟度，但尚未达到“阶段6研究结论已经最终成立”的证据等级。",
        default_bold=True,
    )

    doc.add_paragraph()
    p = doc.add_paragraph(style="Small Note")
    add_inline(
        p,
        r"项目根目录：C:\Users\Lenovo\Desktop\项目\rag-deepspace\uav_spectrum_semcom_project",
    )
    doc.add_page_break()


def add_quick_navigation(doc: Document) -> None:
    heading = doc.add_paragraph("快速导航", style="Heading 1")
    heading.paragraph_format.keep_with_next = True
    rows = [
        ("定位与数据", "第2—5节", "系统是什么、输入如何形成、任务与指标如何定义"),
        ("算法主体", "第6—13节", "阶段4—6演进、码本、codec、状态恢复、预测和全流程"),
        ("证据与边界", "第14—20节", "冻结参数、基线、结果、创新、可行性与负结果"),
        ("审查与复核", "第21—23节", "代码证据、审查问题和当前交接结论"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    for j, text in enumerate(("主题", "位置", "审查内容")):
        cell = table.cell(0, j)
        set_cell_shading(cell, BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, text, default_bold=True)
        for r in p.runs:
            r.font.color.rgb = RGBColor.from_string(WHITE)
    for theme, where, purpose in rows:
        cells = table.add_row().cells
        for j, value in enumerate((theme, where, purpose)):
            set_cell_shading(cells[j], WHITE if len(table.rows) % 2 else LIGHT_GREY)
            p = cells[j].paragraphs[0]
            add_inline(p, value, default_bold=(j == 0))
    widths = [1800, 1300, 6200]
    for row in table.rows:
        for j, width in enumerate(widths):
            set_cell_width(row.cells[j], width)
            set_cell_margins(row.cells[j])
            row.cells[j].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_repeat_table_header(table.rows[0])
    doc.add_paragraph()
    p = doc.add_paragraph(style="Small Note")
    add_inline(p, "Word导航窗格可按标题直接跳转；所有面向读者的公式均采用Unicode可读形式。")


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        content = lines[i].strip().strip("|")
        rows.append([cell.strip() for cell in content.split("|")])
        i += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in rows[1]):
        rows.pop(1)
    return rows, i


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    rows = [row + [""] * (cols - len(row)) for row in rows]
    table = doc.add_table(rows=1, cols=cols)
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.style = "Table Grid"

    lengths = []
    for col in range(cols):
        max_len = max(len(re.sub(r"`|\*\*", "", row[col])) for row in rows)
        lengths.append(max(7.0, min(34.0, max_len ** 0.74 + 4)))
    total = sum(lengths)
    widths = [int(9300 * value / total) for value in lengths]
    widths[-1] += 9300 - sum(widths)

    for j, text in enumerate(rows[0]):
        cell = table.cell(0, j)
        set_cell_shading(cell, BLUE)
        set_cell_width(cell, widths[j])
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        add_inline(p, text, default_bold=True)
        for run in p.runs:
            run.font.color.rgb = RGBColor.from_string(WHITE)
            run.font.size = Pt(8.2 if cols >= 6 else 8.8)
    set_repeat_table_header(table.rows[0])

    for row_idx, values in enumerate(rows[1:], start=1):
        cells = table.add_row().cells
        for j, text in enumerate(values):
            cell = cells[j]
            set_cell_width(cell, widths[j])
            set_cell_margins(cell)
            set_cell_shading(cell, WHITE if row_idx % 2 else LIGHT_GREY)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if j > 0 and re.fullmatch(r"[+\-−]?[0-9.,%—–→\sA-Za-zμ.]+", text):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, text, default_bold=False)
            for run in p.runs:
                run.font.size = Pt(8.0 if cols >= 6 else 8.6)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def insert_figure(doc: Document, path: Path, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.add_run().add_picture(str(path), width=Inches(6.65))
    cap = doc.add_paragraph(style="Small Note")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(cap, caption)


def build_docx() -> None:
    draw_evolution()
    draw_architecture()
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_quick_navigation(doc)
    insert_figure(doc, EVOLUTION, "图1  算法主线从任务表示逐步扩展为有状态可靠闭环")

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    skip_cover_headings = 2
    inserted_architecture = False
    page_break_sections = {
        "8. 阶段6核心算法：解析型任务等价码本",
        "13. 当前冻结系统的逐场景完整流程",
        "16. 阶段6开发结果",
        "17. 已完成工作的创新性判断",
        "20. 证据等级与当前结论",
        "21. 代码、配置与结果索引",
        "22. 复核建议与待审问题",
    }
    formula_prefixes = (
        "xₜ =",
        "M =",
        "c(",
        "j* =",
        "R(",
        "Rₘₐₓ",
        "sₜ(",
        "rₜ,reuse",
        "B_abs",
        "a*(",
        "C(",
    )

    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|"):
            table_rows, i = parse_table(lines, i)
            add_table(doc, table_rows)
            continue
        if stripped.startswith("#"):
            match = re.match(r"^(#{1,3})\s+(.*)$", stripped)
            if match:
                level = len(match.group(1))
                text = match.group(2)
                if skip_cover_headings:
                    skip_cover_headings -= 1
                    i += 1
                    continue
                style = "Heading 1" if level == 2 else "Heading 2" if level == 3 else "Title"
                p = doc.add_paragraph(style=style)
                if level == 2 and text in page_break_sections:
                    p.paragraph_format.page_break_before = True
                add_inline(p, text, default_bold=True)
                if text == "13. 当前冻结系统的逐场景完整流程" and not inserted_architecture:
                    insert_figure(doc, ARCHITECTURE, "图2  阶段6冻结候选的逐场景数据流、控制流与反馈闭环")
                    inserted_architecture = True
                i += 1
                continue
        if stripped.startswith(">"):
            text = stripped[1:].strip()
            table = doc.add_table(rows=1, cols=1)
            table.autofit = False
            set_cell_width(table.cell(0, 0), 9300)
            set_cell_shading(table.cell(0, 0), PALE_BLUE)
            set_cell_margins(table.cell(0, 0), top=125, start=160, bottom=125, end=160)
            p = table.cell(0, 0).paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, text)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            i += 1
            continue
        if re.match(r"^\d+\.\s+", stripped):
            display_number = 1
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i].strip()):
                text = re.sub(r"^\d+\.\s+", "", lines[i].strip())
                add_numbered_item(doc, text, display_number)
                display_number += 1
                i += 1
            continue
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, stripped[2:])
            i += 1
            continue

        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (
                not nxt
                or nxt.startswith("#")
                or nxt.startswith("|")
                or nxt.startswith("> ")
                or nxt.startswith("- ")
                or re.match(r"^\d+\.\s+", nxt)
            ):
                break
            paragraph_lines.append(nxt)
            i += 1
        text = " ".join(paragraph_lines)
        plain = text.strip("`")
        if stripped.startswith("`") and stripped.endswith("`") and any(plain.startswith(prefix) for prefix in formula_prefixes):
            p = doc.add_paragraph(style="Formula Readable")
            add_inline(p, text)
        else:
            p = doc.add_paragraph()
            add_inline(p, text)

    # Ensure every section has the same header/footer settings after page breaks.
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.78)
        section.bottom_margin = Inches(0.72)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
        section.header_distance = Inches(0.32)
        section.footer_distance = Inches(0.32)

    doc.core_properties.title = "基于离散任务导向语义通信的频谱资源选择系统：完整算法、实现与证据审查稿"
    doc.core_properties.subject = "阶段4—6算法与证据第三方审查交接"
    doc.core_properties.author = "项目研究者（Codex辅助整理）"
    doc.core_properties.keywords = "任务导向语义通信, 频谱资源, 任务等价码本, 事件触发, 上下文恢复"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_docx()
